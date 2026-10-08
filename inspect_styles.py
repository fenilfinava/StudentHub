from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document('PRAC_1_ITUE203_25CS011.docx')
for i, p in enumerate(doc.paragraphs[:5]):
    align = p.alignment
    runs = [(r.text, r.bold, r.italic, r.font.size, r.font.name) for r in p.runs]
    print(f"P{i}: Align: {align}, Runs: {runs}")

print("--- Tables ---")
for i, t in enumerate(doc.tables[:3]):
    print(f"Table {i} style: {t.style.name}")

print("Table 1 contents:")
for r in doc.tables[1].rows:
    print([c.text for c in r.cells])

print("Table 2 cell 0,0:")
print(doc.tables[2].cell(0,0).text)

